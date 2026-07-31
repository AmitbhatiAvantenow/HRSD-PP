import base64
import calendar
import json
import re
import uuid
from datetime import datetime
from io import BytesIO

from docx import Document as DocxDocument

from odoo import api, fields, models

from . import document_generation_engine as engine

VARIABLE_TOKEN_RE = re.compile(r'\{\{\s*([a-zA-Z0-9_]+)\s*\}\}')

DEFAULT_CANVAS = '{"page": {}, "blocks": []}'
POPULAR_THRESHOLD = 5
THUMBNAIL_MAX_BLOCKS = 5
THUMBNAIL_TEXT_LEN = 160


def _thumbnail_preview(canvas_data_str):
    """A lightweight, genuinely-live preview of a template's first few blocks (in
    reading order) for the template card thumbnail -- not a rasterized image, just
    the actual heading/text/table content, truncated, so the card always reflects
    the current canvas without any separate rendering/caching step."""
    try:
        canvas = json.loads(canvas_data_str or DEFAULT_CANVAS)
    except (ValueError, TypeError):
        canvas = json.loads(DEFAULT_CANVAS)
    blocks = canvas.get('blocks', [])
    page_count = sum(1 for b in blocks if b.get('type') == 'page_break') + 1
    thumb = []
    for b in sorted(blocks, key=lambda b: (b.get('y', 0), b.get('x', 0))):
        btype = b.get('type')
        props = b.get('props') or {}
        if btype in ('heading', 'text'):
            text = (props.get('text') or '').strip()
            if not text:
                continue
            thumb.append({'type': btype, 'text': text[:THUMBNAIL_TEXT_LEN]})
        elif btype == 'table':
            thumb.append({
                'type': 'table',
                'headers': (props.get('headers') or [])[:3],
                'rows': [row[:3] for row in (props.get('rows') or [])[:2]],
            })
        if len(thumb) >= THUMBNAIL_MAX_BLOCKS:
            break
    return {'blocks': thumb, 'page_count': page_count}


class DocumentTemplate(models.Model):
    _name = 'document.template'
    _inherit = ['mail.thread']
    _description = 'Document Template'
    _order = 'write_date desc'
    _rec_name = 'name'

    name = fields.Char(required=True, tracking=True)
    category_id = fields.Many2one('document.template.category', required=True, tracking=True)
    department = fields.Selection(related='category_id.department', store=True, readonly=True)
    description = fields.Text()
    language = fields.Selection([
        ('en', 'English'), ('ar', 'Arabic'), ('fr', 'French'), ('hi', 'Hindi'),
    ], default='en', required=True)
    paper_size = fields.Selection([
        ('a4', 'A4'), ('letter', 'Letter'), ('legal', 'Legal'),
    ], default='a4', required=True)
    orientation = fields.Selection([
        ('portrait', 'Portrait'), ('landscape', 'Landscape'),
    ], default='portrait', required=True)
    access_level = fields.Selection([
        ('private', 'Private (only me)'), ('team', 'Team (my department)'), ('company', 'Company-wide'),
    ], default='private', required=True)
    tag_ids = fields.Many2many('document.template.tag')

    status = fields.Selection([
        ('draft', 'Draft'), ('published', 'Published'),
    ], default='draft', required=True, tracking=True)
    active = fields.Boolean(default=True)
    source_format = fields.Selection([
        ('blank', 'Built from scratch'), ('docx', 'Uploaded Word Doc'),
    ], default='blank', required=True)

    variable_ids = fields.One2many('document.template.variable', 'template_id')
    canvas_data = fields.Text(default=DEFAULT_CANVAS)

    favorite_user_ids = fields.Many2many(
        'res.users', 'document_template_favorite_rel', 'template_id', 'user_id', string='Favourited By')
    shared_user_ids = fields.Many2many(
        'res.users', 'document_template_shared_rel', 'template_id', 'user_id', string='Shared With')
    rating = fields.Float(default=4.0)

    generated_ids = fields.One2many('document.generated', 'template_id')
    usage_count = fields.Integer(compute='_compute_usage_count', store=True)
    is_popular = fields.Boolean(compute='_compute_usage_count', store=True)

    approval_state = fields.Selection([
        ('none', 'Not Submitted'), ('pending', 'Pending Approval'),
        ('approved', 'Approved'), ('rejected', 'Rejected'),
    ], default='none', tracking=True)
    approver_id = fields.Many2one('res.users')
    approval_note = fields.Text()

    @api.depends('generated_ids')
    def _compute_usage_count(self):
        for rec in self:
            rec.usage_count = len(rec.generated_ids)
            rec.is_popular = rec.usage_count >= POPULAR_THRESHOLD

    # ------------------------------------------------------------------
    # Approval workflow
    # ------------------------------------------------------------------

    def action_submit_for_approval(self):
        self.write({'approval_state': 'pending'})

    def action_approve(self, note=False):
        self.write({'approval_state': 'approved', 'approval_note': note, 'approver_id': self.env.user.id})

    def action_reject(self, note=False):
        self.write({'approval_state': 'rejected', 'approval_note': note, 'approver_id': self.env.user.id})

    def action_toggle_favorite(self):
        self.ensure_one()
        uid = self.env.user.id
        if uid in self.favorite_user_ids.ids:
            self.favorite_user_ids = [(3, uid)]
        else:
            self.favorite_user_ids = [(4, uid)]
        return uid in self.favorite_user_ids.ids

    # ------------------------------------------------------------------
    # Wizard-meta / grid RPCs (OWL frontend)
    # ------------------------------------------------------------------

    @api.model
    def get_template_wizard_meta(self):
        return {
            'categories': self.env['document.template.category'].search_read([], ['name', 'department']),
            'languages': [{'key': k, 'label': v} for k, v in self._fields['language'].selection],
            'paper_sizes': [{'key': k, 'label': v} for k, v in self._fields['paper_size'].selection],
            'orientations': [{'key': k, 'label': v} for k, v in self._fields['orientation'].selection],
            'access_levels': [{'key': k, 'label': v} for k, v in self._fields['access_level'].selection],
            'tags': self.env['document.template.tag'].search_read([], ['name']),
        }

    @api.model
    def create_template_wizard(self, vals):
        if not (vals.get('name') or '').strip():
            raise ValueError('Template name is required.')
        if not vals.get('category_id'):
            raise ValueError('Category is required.')
        template = self.create({
            'name': vals['name'].strip(),
            'category_id': vals['category_id'],
            'description': vals.get('description') or '',
            'language': vals.get('language') or 'en',
            'paper_size': vals.get('paper_size') or 'a4',
            'orientation': vals.get('orientation') or 'portrait',
            'access_level': vals.get('access_level') or 'private',
            'tag_ids': [(6, 0, vals.get('tag_ids') or [])],
        })
        return {'template_id': template.id}

    @api.model
    def create_from_upload(self, vals, file_data, file_name):
        if not (vals.get('name') or '').strip():
            raise ValueError('Template name is required.')
        if not vals.get('category_id'):
            raise ValueError('Category is required.')
        if not (file_name or '').lower().endswith('.docx'):
            raise ValueError('Only Word (.docx) files can be uploaded and made editable.')
        template = self.create({
            'name': vals['name'].strip(),
            'category_id': vals['category_id'],
            'description': vals.get('description') or '',
            'language': vals.get('language') or 'en',
            'paper_size': vals.get('paper_size') or 'a4',
            'orientation': vals.get('orientation') or 'portrait',
            'access_level': vals.get('access_level') or 'private',
            'tag_ids': [(6, 0, vals.get('tag_ids') or [])],
            'source_format': 'docx',
        })
        template._import_docx(file_data)
        return {'template_id': template.id}

    def _import_docx(self, file_data):
        """Parse an uploaded .docx into editable text/heading blocks, stacked top to
        bottom in reading order, and auto-register any {{ variable }} tokens already
        present in the text as document.template.variable records."""
        self.ensure_one()
        docx_file = DocxDocument(BytesIO(base64.b64decode(file_data)))

        margin = 40
        block_w = 500
        _, page_h = engine.page_dims(self.paper_size, self.orientation)
        page_bottom = page_h - margin  # continuous-coordinate y where page 1 runs out of room

        blocks = []
        y = margin
        found_keys = []
        for para in docx_file.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style_name = (para.style.name or '') if para.style else ''
            is_heading = style_name.lower().startswith('heading') or style_name.lower() == 'title'
            level_match = re.search(r'(\d+)', style_name)
            level = int(level_match.group(1)) if level_match and is_heading else 1
            bold = is_heading or any(r.bold for r in para.runs if r.bold)
            props = {
                'text': text,
                'font_size': 20 if is_heading else 11,
                'bold': bold,
                'italic': False,
                'align': 'left',
                'color': '#111111',
            }
            if is_heading:
                props['level'] = level
            block_type = 'heading' if is_heading else 'text'
            # Measure the REAL wrapped height with the same Paragraph/style logic the
            # PDF renderer draws with, so a block never overlaps the one after it --
            # a naive "count newlines" estimate badly under-measures long paragraphs
            # that wrap across many lines at this width.
            h = max(24, engine.measure_text_height(text, props, block_type, block_w) + 8)

            if y + h > page_bottom:
                blocks.append({
                    'id': f'b_{uuid.uuid4().hex[:8]}', 'type': 'page_break',
                    'x': 0, 'y': page_bottom, 'w': 0, 'h': 0, 'z': 0, 'props': {},
                })
                y = page_bottom + margin
                page_bottom += page_h

            blocks.append({
                'id': f'b_{uuid.uuid4().hex[:8]}',
                'type': block_type,
                'x': margin, 'y': y, 'w': block_w, 'h': h, 'z': 1,
                'props': props,
            })
            y += h + 12
            found_keys.extend(VARIABLE_TOKEN_RE.findall(text))

        self.write({'canvas_data': json.dumps({'page': {}, 'blocks': blocks})})

        existing_keys = set(self.variable_ids.mapped('key'))
        seq = (max(self.variable_ids.mapped('sequence'), default=0) or 0) + 10
        Variable = self.env['document.template.variable']
        for key in dict.fromkeys(found_keys):  # de-dupe, preserve first-seen order
            if key in existing_keys:
                continue
            Variable.create({
                'template_id': self.id,
                'name': key.replace('_', ' ').title(),
                'key': key,
                'sequence': seq,
            })
            existing_keys.add(key)
            seq += 10

    @api.model
    def get_grid_data(self, domain=None, search=None, category_id=None, order=None):
        domain = list(domain or [])
        if search:
            domain += [('name', 'ilike', search)]
        if category_id:
            domain += [('category_id', '=', int(category_id))]
        order_map = {
            'newest': 'write_date desc', 'name': 'name asc',
            'rating': 'rating desc', 'popular': 'usage_count desc',
        }
        uid = self.env.user.id
        templates = self.search(domain, order=order_map.get(order, 'write_date desc'))
        return [{
            'id': t.id,
            'name': t.name,
            'category_name': t.category_id.name,
            'category_icon': t.category_id.icon,
            'department': t.department,
            'rating': t.rating,
            'is_popular': t.is_popular,
            'usage_count': t.usage_count,
            'status': t.status,
            'access_level': t.access_level,
            'approval_state': t.approval_state,
            'is_favorite': uid in t.favorite_user_ids.ids,
            'write_date': fields.Datetime.to_string(t.write_date),
            'updated_by': t.write_uid.name,
            'source_format': t.source_format,
            'thumbnail': _thumbnail_preview(t.canvas_data),
        } for t in templates]

    # ------------------------------------------------------------------
    # Builder RPCs
    # ------------------------------------------------------------------

    @api.model
    def get_builder_data(self, template_id):
        template = self.browse(template_id)
        page_w, page_h = engine.page_dims(template.paper_size, template.orientation)
        try:
            canvas = json.loads(template.canvas_data or DEFAULT_CANVAS)
        except (ValueError, TypeError):
            canvas = json.loads(DEFAULT_CANVAS)
        return {
            'template': {
                'id': template.id,
                'name': template.name,
                'paper_size': template.paper_size,
                'orientation': template.orientation,
            },
            'page_width_pt': page_w,
            'page_height_pt': page_h,
            'canvas': canvas,
            'variables': template.variable_ids.read(['name', 'key', 'variable_type', 'default_value', 'is_required']),
        }

    @api.model
    def save_canvas(self, template_id, canvas_json):
        template = self.browse(template_id)
        json.loads(canvas_json)  # validate before persisting
        template.write({'canvas_data': canvas_json})
        return True

    def rename_template(self, name):
        self.ensure_one()
        if (name or '').strip():
            self.write({'name': name.strip()})

    # ------------------------------------------------------------------
    # Generation
    # ------------------------------------------------------------------

    @api.model
    def get_generate_wizard_data(self, template_id):
        template = self.browse(template_id)
        return {
            'template_name': template.name,
            'variables': template.variable_ids.read(['name', 'key', 'variable_type', 'default_value', 'is_required']),
        }

    def _render_bytes(self, variable_values, doc_format):
        self.ensure_one()
        try:
            canvas = json.loads(self.canvas_data or DEFAULT_CANVAS)
        except (ValueError, TypeError):
            canvas = json.loads(DEFAULT_CANVAS)
        variables = engine.coerce_variable_values(self, variable_values)
        page_w, page_h = engine.page_dims(self.paper_size, self.orientation)
        blocks = [engine.substitute_block_props(b, variables) for b in canvas.get('blocks', [])]
        pages = engine.partition_pages(blocks, page_h)
        page_cfg = canvas.get('page') or {}
        page_meta = {
            **page_cfg,
            'width_pt': page_w,
            'height_pt': page_h,
            'header_text': engine.render_text(page_cfg.get('header_text', ''), variables),
            'footer_text': engine.render_text(page_cfg.get('footer_text', ''), variables),
        }
        if doc_format == 'pdf':
            return engine.render_pdf(page_meta, pages)
        if doc_format == 'docx':
            return engine.render_docx(page_meta, pages)
        raise ValueError("doc_format must be 'pdf' or 'docx'")

    def preview_pdf_base64(self, variable_values=None):
        self.ensure_one()
        values = dict(variable_values or {})
        for v in self.variable_ids:
            if v.key not in values or values[v.key] in (None, ''):
                values[v.key] = v.default_value or f'[{v.name}]'
                # a required check would reject placeholders, so bypass it for preview
        try:
            pdf_bytes = self._render_bytes(values, 'pdf')
        except Exception:
            # Preview should never hard-fail the builder; fall back to an empty doc.
            pdf_bytes = engine.render_pdf({'width_pt': 595, 'height_pt': 842}, [[]])
        return base64.b64encode(pdf_bytes).decode()

    def action_generate(self, variable_values, formats, partner_id=False):
        self.ensure_one()
        vals = {
            'template_id': self.id,
            'partner_id': partner_id or False,
            'variable_values': json.dumps(variable_values or {}),
        }
        if 'pdf' in (formats or []):
            vals['file_data_pdf'] = base64.b64encode(self._render_bytes(variable_values, 'pdf'))
            vals['file_name_pdf'] = f'{self.name}.pdf'
        if 'docx' in (formats or []):
            vals['file_data_docx'] = base64.b64encode(self._render_bytes(variable_values, 'docx'))
            vals['file_name_docx'] = f'{self.name}.docx'
        doc = self.env['document.generated'].create(vals)
        return {
            'generated_id': doc.id,
            'download_pdf_url': f'/web/content/document.generated/{doc.id}/file_data_pdf/{doc.file_name_pdf}?download=true' if vals.get('file_data_pdf') else False,
            'download_docx_url': f'/web/content/document.generated/{doc.id}/file_data_docx/{doc.file_name_docx}?download=true' if vals.get('file_data_docx') else False,
        }

    # ------------------------------------------------------------------
    # Dashboard
    # ------------------------------------------------------------------

    @api.model
    def get_dashboard_data(self):
        Template = self.env['document.template']
        Generated = self.env['document.generated']
        uid = self.env.user.id

        month_start = fields.Datetime.now().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        stats = {
            'total_templates': Template.search_count([]),
            'generated_this_month': Generated.search_count([('generated_date', '>=', fields.Datetime.to_string(month_start))]),
            'shared_templates': Template.search_count([('access_level', '!=', 'private')]),
            'favourite_templates': Template.search_count([('favorite_user_ids', 'in', [uid])]),
            'draft_templates': Template.search_count([('status', '=', 'draft')]),
            'published_templates': Template.search_count([('status', '=', 'published')]),
        }

        most_used = Template.search_read([('usage_count', '>', 0)], ['name', 'usage_count'], order='usage_count desc', limit=8)

        by_department = []
        dept_counts = dict(Template._read_group([], ['department'], ['__count']))
        department_selection = Template.fields_get(['department'])['department']['selection']
        for key, label in department_selection:
            count = dept_counts.get(key, 0)
            if count:
                by_department.append({'department': key, 'label': label, 'count': count})

        generated_trend = []
        today = datetime.now()
        y, m = today.year, today.month
        months = []
        for _i in range(6):
            months.append((y, m))
            m -= 1
            if m == 0:
                m = 12
                y -= 1
        months.reverse()
        for (yy, mm) in months:
            last_day = calendar.monthrange(yy, mm)[1]
            start = datetime(yy, mm, 1)
            end = datetime(yy, mm, last_day, 23, 59, 59)
            count = Generated.search_count([
                ('generated_date', '>=', fields.Datetime.to_string(start)),
                ('generated_date', '<=', fields.Datetime.to_string(end)),
            ])
            generated_trend.append({'month_label': start.strftime('%b'), 'count': count})

        approval_breakdown = []
        approval_counts = dict(Template._read_group([], ['approval_state'], ['__count']))
        for key, label in Template._fields['approval_state'].selection:
            approval_breakdown.append({'state': key, 'label': label, 'count': approval_counts.get(key, 0)})

        categories = self.env['document.template.category'].search([])
        cat_counts = dict(Template._read_group([], ['category_id'], ['__count']))
        total_cat = sum(cat_counts.values()) or 1
        category_breakdown = []
        cumulative = 0.0
        for cat in categories:
            count = cat_counts.get(cat, 0)
            if not count:
                continue
            pct = (count / total_cat) * 100
            category_breakdown.append({
                'name': cat.name, 'count': count, 'pct': pct, 'cumulative_pct': cumulative,
            })
            cumulative += pct

        recently_modified = Template.search_read([], ['name', 'write_date', 'status'], order='write_date desc', limit=8)

        return {
            'stats': stats,
            'most_used': most_used,
            'by_department': by_department,
            'generated_trend': generated_trend,
            'approval_breakdown': approval_breakdown,
            'category_breakdown': category_breakdown,
            'recently_modified': recently_modified,
        }

    # ------------------------------------------------------------------
    # Settings (lightweight, backed by ir.config_parameter)
    # ------------------------------------------------------------------

    @api.model
    def get_settings(self):
        ICP = self.env['ir.config_parameter'].sudo()
        return {
            'default_access_level': ICP.get_param('document_templates.default_access_level', 'private'),
            'require_approval': ICP.get_param('document_templates.require_approval', 'False') == 'True',
        }

    @api.model
    def set_settings(self, vals):
        ICP = self.env['ir.config_parameter'].sudo()
        if 'default_access_level' in vals:
            ICP.set_param('document_templates.default_access_level', vals['default_access_level'])
        if 'require_approval' in vals:
            ICP.set_param('document_templates.require_approval', str(bool(vals['require_approval'])))
        return True
